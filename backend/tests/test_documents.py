import io

import docx
import pytest

from modeldock import documents


def test_extract_plain_text():
    result = documents.extract("你好，短剧旁白。".encode("utf-8"), "note.txt")
    assert result["text"] == "你好，短剧旁白。"
    assert result["truncated"] is False


def test_extract_docx():
    document = docx.Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    buffer = io.BytesIO()
    document.save(buffer)
    result = documents.extract(buffer.getvalue(), "script.docx")
    assert "第一段" in result["text"] and "第二段" in result["text"]


def test_reject_binary():
    with pytest.raises(documents.DocumentError):
        documents.extract(b"\x00\x01\x02", "file.bin")
