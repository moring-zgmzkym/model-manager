"""Turn uploaded documents into plain text for chat context.

The browser cannot read PDF/DOCX contents reliably, so ChatPanel posts the
file to /api/extract and receives text back.  Text files are decoded with a
few common encodings; PDF and DOCX use optional libraries that are part of
requirements.txt.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

MAX_CHARS = 200_000
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css", ".log", ".ini",
    ".toml", ".xml", ".sql", ".bat", ".cmd", ".sh", ".ps1", ".java", ".c", ".cpp",
    ".h", ".hpp", ".go", ".rs", ".rb", ".php", ".vue", ".svelte", ".srt", ".vtt",
}


class DocumentError(Exception):
    pass


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise DocumentError("读取 PDF 需要安装 pypdf：pip install pypdf") from exc
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)


def _docx_text(data: bytes) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise DocumentError("读取 Word 需要安装 python-docx：pip install python-docx") from exc
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract(data: bytes, filename: str) -> dict[str, Any]:
    """Return {name, text, chars, truncated} for one uploaded file."""
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(data)
    elif suffix == ".docx":
        text = _docx_text(data)
    elif suffix == ".doc":
        raise DocumentError("暂不支持旧版 .doc，请另存为 .docx 或 PDF")
    elif suffix in TEXT_SUFFIXES or not suffix:
        text = _decode_text(data)
    else:
        if b"\x00" in data[:4096]:
            raise DocumentError(f"暂不支持读取 {suffix or '这种'} 文件")
        text = _decode_text(data)

    text = text.strip()
    if not text:
        raise DocumentError("没有从文件里读到文字")
    truncated = len(text) > MAX_CHARS
    text = text[:MAX_CHARS]
    return {"name": filename, "text": text, "chars": len(text), "truncated": truncated}
